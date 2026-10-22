# -*- coding: utf-8 -*-
"""检查桌面上的使用说明 Word 文档内容是否完整、无乱码"""
import io
import sys
from docx import Document

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
p = r"C:\Users\27748\Desktop\安静陪伴-使用说明.docx"
d = Document(p)
ps = [x.text for x in d.paragraphs if x.text.strip()]
alltext = "\n".join(ps) + "\n" + "\n".join(
    c.text for t in d.tables for r in t.rows for c in r.cells)

print("段落 %d / 表格 %d / 全文字符 %d" % (len(ps), len(d.tables), len(alltext)))
checks = [
    "路径 0：APK", "路径 1：PWA", "路径 2：自己的服务器", "两种连接方式的区别",
    "手机直连（默认）", "电脑关机能用吗", "启动手机版(公网隧道).bat",
    "添加到主屏幕", "安装未知来源", "危机干预热线",
    "导出全部数据为 JSON", "Render", "deepseek-flash",
]
for c in checks:
    print(("  OK   " if c in alltext else "  缺失 ") + c)
print("乱码字符数:", alltext.count("\ufffd"))
