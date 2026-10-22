# -*- coding: utf-8 -*-
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from docx import Document

p = r"C:\Users\27748\Desktop\项目资源_Projects\项目详细流程\V3融合多元先验与几何约束的古建筑图像稠密匹配.docx"
d = Document(p)
print("=== 段落 ===")
for i, para in enumerate(d.paragraphs):
    t = para.text.strip()
    if t:
        print("[%03d][%s] %s" % (i, para.style.name, t))
print("\n=== 表格 (%d) ===" % len(d.tables))
for ti, tb in enumerate(d.tables):
    print("--- table %d (%d行 x %d列) ---" % (ti, len(tb.rows), len(tb.columns)))
    for r in tb.rows:
        cells = [c.text.strip().replace("\n", " / ") for c in r.cells]
        print("  | " + " | ".join(cells))
print("\n=== 内嵌图片 ===")
for rel in d.part.rels.values():
    if "image" in rel.reltype:
        print(" ", rel.rId, rel.target_ref)
